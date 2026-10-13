"""
PDF → Word that keeps the page: each PDF page becomes a Word section of the same size, and everything on it is
rebuilt from what the PDF itself says — the text with its fonts and positions, the rules and boxes it draws,
its tables and its pictures — or, for a scanned page, from OCR lines plus the visual elements cropped from the
scan. Positions become native Word constructs (exact line pitch and space before, indents, tab stops, tables,
paragraph borders, section columns), never runs of spaces or empty lines.

Libraries: pdfplumber/pdfminer.six (MIT) to read the page, pypdfium2 to render, pikepdf to pull embedded
images, python-docx (MIT) to write. Coordinates are PDF points, top-left origin, as pdfplumber gives them.

    build_docx(pdf_path, out_path, pages=[{"index": 0, "ocr": None | [line...]}, ...])

An OCR line is {"text", "box": [x0, top, x1, bottom] (points), "size"?, "bold"?}.
"""

from __future__ import annotations

import io
import logging
import math
import re
import statistics
from dataclasses import dataclass, field

import numpy as np
from PIL import Image

log = logging.getLogger("image-service")

EMU_PER_PT = 12700
TWIPS = 20  # per point
# Word's "exactly" line box puts the baseline this far down the box (measured on Word and LibreOffice renders).
BASELINE_RATIO = 0.8
DEVANAGARI = re.compile("[ऀ-ॿ]")
BROKEN_TEXT = re.compile("[\x00�]|(?:^|\\s)[ा-्ँ-ः]")


# ------------------------------------------------------------------------------------------------ model


@dataclass
class Word:
    text: str
    x0: float
    x1: float
    top: float
    bottom: float
    size: float
    font: str
    bold: bool
    italic: bool
    colour: str | None
    underline: bool = False
    ocr: bool = False  # read by OCR: its size is an estimate


@dataclass
class Line:
    words: list[Word]

    @property
    def x0(self):
        return min(w.x0 for w in self.words)

    @property
    def x1(self):
        return max(w.x1 for w in self.words)

    @property
    def top(self):
        return min(w.top for w in self.words)

    @property
    def bottom(self):
        return max(w.bottom for w in self.words)

    @property
    def size(self):
        return statistics.median(w.size for w in self.words)

    @property
    def baseline(self):
        # pdfminer's glyph box runs from about 0.8 em above the baseline to 0.2 em below it.
        return self.bottom - 0.2 * self.size

    @property
    def text(self):
        return " ".join(w.text for w in self.words)


@dataclass
class Element:
    """Something placed on the page: a text block, table, picture or rule. Box in points."""

    kind: str  # "text" | "table" | "image" | "rule"
    x0: float
    top: float
    x1: float
    bottom: float
    lines: list[Line] = field(default_factory=list)
    data: dict = field(default_factory=dict)

    @property
    def width(self):
        return self.x1 - self.x0

    @property
    def height(self):
        return self.bottom - self.top


@dataclass
class PageModel:
    width: float
    height: float
    elements: list[Element]
    shading: list[tuple[tuple[float, float, float, float], str]]  # filled boxes behind text
    floating: list[Element]  # pictures drawn under the text (backgrounds, watermarks)


# ------------------------------------------------------------------------------------------------ reading


def colour_hex(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (int, float)):
        value = (value,)
    value = tuple(float(v) for v in value if isinstance(v, (int, float)))
    if len(value) == 1:
        r = g = b = value[0]
    elif len(value) == 3:
        r, g, b = value
    elif len(value) == 4:
        c, m, y, k = value
        r, g, b = (1 - c) * (1 - k), (1 - m) * (1 - k), (1 - y) * (1 - k)
    else:
        return None
    return "%02X%02X%02X" % tuple(max(0, min(255, round(v * 255))) for v in (r, g, b))


STYLE_BOLD = {"bold", "bd", "black", "blk", "heavy", "hv", "semibold", "sb", "demi", "demibold", "medi", "medium", "md", "extrabold", "xbold", "ultra"}
STYLE_ITALIC = {"italic", "ital", "it", "oblique", "obl", "slant", "sl", "inclined"}


def font_info(fontname: str) -> tuple[str, bool, bool]:
    """A PDF font name → (Word family, bold, italic). Subset prefixes go; metric-compatible families are used."""
    name = fontname.split("+", 1)[-1]
    low = name.lower().replace(" ", "")
    base, _, style = name.partition("-")
    if not style:
        style = name.partition(",")[2]
    # Style words, from the part after the hyphen ("Bd", "BdIt", "ReguItal", "MediItal", "BoldOblique").
    tokens = {t.lower() for t in re.findall(r"[A-Z][a-z]*|[a-z]+", style)}
    bold = bool(tokens & STYLE_BOLD) or bool(re.search(r"bold|black|heavy|semibold|demi", low))
    italic = bool(tokens & STYLE_ITALIC) or bool(re.search(r"italic|oblique", low))
    # TeX's Computer Modern: cmbx = bold, cmti/cmsl = italic/slanted.
    if re.match(r"(cm|lm)bx", low):
        bold = True
    if re.match(r"(cm|lm)(ti|sl|mi)", low):
        italic = True
    families = [
        (r"helvetica|arial|liberationsans|arimo|nimbussan|texgyreheros|franklingothic", "Arial"),
        (r"timesnewroman|times|liberationserif|tinos|nimbusrom|texgyretermes|^cmr|^cmbx|^cmti|^cmsl|^lmroman|^sfrm|^ptmr", "Times New Roman"),
        (r"courier|liberationmono|cousine|nimbusmon|^cmtt|^lmmono", "Courier New"),
        (r"georgia", "Georgia"),
        (r"calibri|carlito", "Calibri"),
        (r"cambria|caladea", "Cambria"),
        (r"verdana", "Verdana"),
        (r"tahoma", "Tahoma"),
        (r"garamond", "Garamond"),
        (r"segoe", "Segoe UI"),
        (r"consolas", "Consolas"),
        (r"^cmsy|^cmmi|^cmex|^msbm|^msam|symbol", "Cambria Math"),
        (r"mangal|kalimati|kokila|aparajita|utsaah|nirmala|notosansdevanagari|kohinoordevanagari|devanagarimt|preeti", None),
    ]
    for pattern, family in families:
        if re.search(pattern, low):
            if family is None:
                break  # a Devanagari face: keep its own name (Word falls back to a Devanagari font if it's missing)
            return family, bold, italic
    # CamelCase → words ("KohinoorDevanagari" → "Kohinoor Devanagari").
    base = re.sub(r"(PSMT|PS|MT|Std|Pro|LT)$", "", base)
    pretty = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", base).strip() or "Arial"
    return pretty, bold, italic


CONTROL = re.compile("[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")


def read_words(page) -> tuple[list[Word], list[tuple]]:
    """The page's words with their look; and the boxes of rotated text (kept as part of a picture)."""
    raw = page.extract_words(x_tolerance=1.5, y_tolerance=2, keep_blank_chars=False, use_text_flow=False, extra_attrs=["fontname", "size", "non_stroking_color", "upright"])
    words, rotated = [], []
    for w in raw:
        if not w.get("upright", True):
            rotated.append((w["x0"], w["top"], w["x1"], w["bottom"]))
            continue
        family, bold, italic = font_info(w["fontname"])
        words.append(Word(w["text"], w["x0"], w["x1"], w["top"], w["bottom"], float(w["size"]), family, bold, italic, colour_hex(w.get("non_stroking_color"))))
    return words, rotated


def group_lines(words: list[Word], split: float = 1.6) -> list[Line]:
    """Words → lines: same row (overlapping vertically), then split where a gap is too wide to be a space."""
    rows: list[list[Word]] = []
    for word in sorted(words, key=lambda w: (round(w.bottom), w.x0)):
        middle = (word.top + word.bottom) / 2
        for row in rows:
            ref = row[0]
            if abs((ref.top + ref.bottom) / 2 - middle) < 0.45 * max(ref.bottom - ref.top, word.bottom - word.top):
                row.append(word)
                break
        else:
            rows.append([word])
    lines: list[Line] = []
    for row in rows:
        row.sort(key=lambda w: w.x0)
        current = [row[0]]
        for word in row[1:]:
            gap = word.x0 - current[-1].x1
            size = max(word.size, current[-1].size)
            # Wider than ~two spaces of justified text: a new chunk (a tab stop, a column, a cell).
            if gap > max(split * size, 14 if split >= 1 else 3):
                lines.append(Line(current))
                current = [word]
            else:
                current.append(word)
        lines.append(Line(current))
    return lines


def segments(page) -> tuple[list[dict], list[dict]]:
    """The page's drawn rules (thin lines, thin rectangles and thin paths) and filled boxes."""
    rules, boxes = [], []
    shapes = [*page.lines, *page.rects, *page.curves]
    for shape in shapes:
        x0, x1, top, bottom = shape["x0"], shape["x1"], shape["top"], shape["bottom"]
        w, h = x1 - x0, bottom - top
        filled = shape.get("fill")
        colour = colour_hex(shape.get("non_stroking_color") if filled else shape.get("stroking_color")) or "000000"
        thick = shape.get("linewidth") or 0.75
        if h <= 2.5 and w > 1:
            rules.append({"dir": "h", "x0": x0, "x1": x1, "y": (top + bottom) / 2, "thick": max(h, thick if not filled else h, 0.5), "colour": colour})
        elif w <= 2.5 and h > 1:
            rules.append({"dir": "v", "y0": top, "y1": bottom, "x": (x0 + x1) / 2, "thick": max(w, 0.5), "colour": colour})
        elif shape.get("object_type") == "rect":
            if filled and colour not in ("FFFFFF",):
                boxes.append({"box": (x0, top, x1, bottom), "fill": colour})
            if shape.get("stroke"):
                stroke = colour_hex(shape.get("stroking_color")) or "000000"
                for side in (("h", x0, x1, top), ("h", x0, x1, bottom)):
                    rules.append({"dir": "h", "x0": side[1], "x1": side[2], "y": side[3], "thick": thick, "colour": stroke})
                for x in (x0, x1):
                    rules.append({"dir": "v", "y0": top, "y1": bottom, "x": x, "thick": thick, "colour": stroke})
    return merge_rules(rules), boxes


def merge_rules(rules: list[dict]) -> list[dict]:
    """Collinear pieces → one rule; many short pieces with gaps → one dashed rule."""
    out = []
    for direction in ("h", "v"):
        key, a, b = ("y", "x0", "x1") if direction == "h" else ("x", "y0", "y1")
        items = sorted((r for r in rules if r["dir"] == direction), key=lambda r: (round(r[key] * 2), r[a]))
        current = None
        for rule in items:
            if current and abs(current[key] - rule[key]) <= 1.0 and rule[a] - current[b] <= 6:
                gap = rule[a] - current[b]
                if gap > 0.5:
                    current["gaps"] += 1
                current[b] = max(current[b], rule[b])
                current["pieces"] += 1
                current["thick"] = max(current["thick"], rule["thick"])
            else:
                if current:
                    out.append(current)
                current = {**rule, "gaps": 0, "pieces": 1}
        if current:
            out.append(current)
    for rule in out:
        rule["dashed"] = rule["pieces"] >= 6 and rule["gaps"] >= rule["pieces"] // 2
    return out


# ------------------------------------------------------------------------------------------------ pictures


def embedded_image(pdf_pike, page_index: int, name: str | None) -> Image.Image | None:
    """The picture as stored in the PDF (with its transparency), or None to fall back to a render."""
    if pdf_pike is None or not name:
        return None
    try:
        import pikepdf
        from pikepdf import PdfImage

        page = pdf_pike.pages[page_index]
        xobjects = page.resources.get("/XObject", {})
        obj = xobjects.get("/" + name.lstrip("/"))
        if obj is None:
            return None
        image = PdfImage(obj).as_pil_image()
        smask = obj.get("/SMask")
        if smask is not None:
            mask = PdfImage(smask).as_pil_image().convert("L")
            if mask.size != image.size:
                mask = mask.resize(image.size)
            image = image.convert("RGB")
            image.putalpha(mask)
        elif image.mode not in ("RGB", "RGBA", "L"):
            image = image.convert("RGB")
        return image
    except Exception:  # noqa: BLE001
        return None


_last_render: dict = {}


def render_region(pdf_doc, page_index: int, box: tuple[float, float, float, float], dpi: int = 200) -> Image.Image:
    scale = dpi / 72
    # A page with several pictures is rendered once (only the latest page is kept, so memory stays flat).
    key = (id(pdf_doc), page_index, dpi)
    if _last_render.get("key") != key:
        _last_render.clear()
        _last_render.update(key=key, image=pdf_doc[page_index].render(scale=scale).to_pil())
    full = _last_render["image"]
    x0, top, x1, bottom = box
    crop = (max(0, int(x0 * scale)), max(0, int(top * scale)), min(full.width, math.ceil(x1 * scale)), min(full.height, math.ceil(bottom * scale)))
    return full.crop(crop).convert("RGB")


def png_bytes(image: Image.Image, max_side: int = 2000) -> bytes:
    if max(image.size) > max_side:
        image = image.copy()
        image.thumbnail((max_side, max_side))
    buffer = io.BytesIO()
    if image.mode == "RGBA" or (image.mode == "L"):
        image.save(buffer, format="PNG", optimize=True)
    else:
        image.convert("RGB").save(buffer, format="JPEG", quality=88)
    return buffer.getvalue()


# ------------------------------------------------------------------------------------------------ page analysis


def overlaps(a, b, pad: float = 0.0) -> bool:
    return a[0] < b[2] + pad and b[0] < a[2] + pad and a[1] < b[3] + pad and b[1] < a[3] + pad


def inside(inner, outer, pad: float = 1.0) -> bool:
    return inner[0] >= outer[0] - pad and inner[1] >= outer[1] - pad and inner[2] <= outer[2] + pad and inner[3] <= outer[3] + pad


def word_box(w: Word):
    return (w.x0, w.top, w.x1, w.bottom)


def analyse_digital(plumber_page, page_index: int, pdf_doc, pdf_pike, ocr_lines: list[dict] | None) -> PageModel:
    page = plumber_page
    width, height = float(page.width), float(page.height)
    words, rotated = read_words(page)
    rules, boxes = segments(page)
    elements: list[Element] = []
    floating: list[Element] = []
    taken = set()  # words placed by a table or a picture

    # Pictures: embedded images (as stored), each with its place on the page.
    for image in page.images:
        box = (max(0, image["x0"]), max(0, image["top"]), min(width, image["x1"]), min(height, image["bottom"]))
        if box[2] - box[0] < 4 or box[3] - box[1] < 4:
            continue
        covered = [i for i, w in enumerate(words) if overlaps(word_box(w), box, -1)]
        area = (box[2] - box[0]) * (box[3] - box[1])
        picture = embedded_image(pdf_pike, page_index, image.get("name")) or render_region(pdf_doc, page_index, box)
        element = Element("image", *box, data={"png": png_bytes(picture)})
        # A picture under most of the page, or under text, is a background: it floats behind the text.
        if area > 0.5 * width * height or len(covered) > 3:
            floating.append(element)
        else:
            elements.append(element)

    # Tables: ruled grids.
    tables = []
    try:
        found = page.find_tables({"vertical_strategy": "lines", "horizontal_strategy": "lines", "snap_tolerance": 3, "join_tolerance": 3, "intersection_tolerance": 3})
    except Exception:  # noqa: BLE001
        found = []
    for table in found:
        x0, top, x1, bottom = table.bbox
        cells = [c for c in table.cells if c[2] - c[0] > 2 and c[3] - c[1] > 2]
        if len(cells) < 1 or (x1 - x0) < 20 or (bottom - top) < 8:
            continue
        tables.append(table)
        mine = [i for i, w in enumerate(words) if inside(word_box(w), table.bbox, 2)]
        taken.update(mine)
        elements.append(Element("table", x0, top, x1, bottom, data={"cells": cells, "words": [words[i] for i in mine]}))

    # Boxes: a rectangle drawn around text (four rules or a stroked frame) is a one-cell table, borders and all.
    table_boxes = [t.bbox for t in tables]
    for box in find_boxes(rules):
        if any(overlaps(box, t, 1) for t in table_boxes):
            continue
        mine = [i for i, w in enumerate(words) if inside(word_box(w), box, 1) and i not in taken]
        if not mine:
            continue
        taken.update(mine)
        table_boxes.append(box)
        elements.append(Element("table", *box, data={"cells": [box], "words": [words[i] for i in mine]}))

    # Vector drawings that aren't rules, boxes or tables (logos, charts, icons, rotated captions): rendered as
    # pictures, with any words inside them, so nothing is drawn twice and nothing is lost.
    drawn = list(rotated)
    for shape in [*page.curves, *page.lines, *page.rects]:
        box = (shape["x0"], shape["top"], shape["x1"], shape["bottom"])
        if any(inside(box, t, 2) for t in table_boxes):
            continue
        thin_h, thin_v = box[3] - box[1] <= 2.5, box[2] - box[0] <= 2.5
        if (thin_h and box[2] - box[0] > 0.25 * width) or (thin_v and box[3] - box[1] > 0.25 * height):
            continue  # a long rule: a divider or a frame edge, not a drawing
        if shape.get("object_type") == "rect" and not thin_h and not thin_v and shape.get("fill") and not shape.get("stroke"):
            continue  # a filled box: shading behind text, handled with the text
        if max(box[2] - box[0], box[3] - box[1]) < 9:
            continue  # bullet dots, chart markers on their own, specks
        drawn.append(box)
    clusters = []
    for cluster in cluster_boxes(drawn, 8):
        members = [d for d in drawn if inside(d, cluster, 0.5)]
        cw, ch = cluster[2] - cluster[0], cluster[3] - cluster[1]
        has_rotated = any(inside(rb, cluster, 0.5) for rb in rotated)
        if not has_rotated and (len(members) < 3 or cw * ch < 400 or min(cw, ch) < 12):
            continue  # a stray tick or dot
        clusters.append(cluster)
    for cluster in clusters:
        x0, top, x1, bottom = cluster
        # Labels just outside the drawing (axis titles) belong to it too.
        grown = (x0 - 2, top - 2, x1 + 2, bottom + 2)
        mine = [i for i, w in enumerate(words) if i not in taken and inside(word_box(w), grown, 4)]
        taken.update(mine)
        if mine:
            x0, top = min(x0, *(words[i].x0 for i in mine)), min(top, *(words[i].top for i in mine))
            x1, bottom = max(x1, *(words[i].x1 for i in mine)), max(bottom, *(words[i].bottom for i in mine))
        box = (max(0, x0 - 1), max(0, top - 1), min(width, x1 + 1), min(height, bottom + 1))
        elements.append(Element("image", *box, data={"png": png_bytes(render_region(pdf_doc, page_index, box)), "words": len(mine)}))

    # Text: lines, with bullets drawn as dots turned back into "•", underlines found from rules just below.
    free = [w for i, w in enumerate(words) if i not in taken and not any(inside(word_box(w), (e.x0, e.top, e.x1, e.bottom), 1) for e in elements if e.kind == "image")]
    if ocr_lines:
        free = repair_with_ocr(free, ocr_lines)
    def is_bullet(shape):
        w, h = shape["x1"] - shape["x0"], shape["bottom"] - shape["top"]
        if not (1.8 <= w <= 8 and 1.8 <= h <= 8 and 0.6 <= w / h <= 1.7):
            return False
        middle = (shape["top"] + shape["bottom"]) / 2
        # A marker sits just left of the text it introduces.
        return any(0 < word.x0 - shape["x1"] < 22 and word.top <= middle <= word.bottom for word in free)

    dots = [c for c in [*page.curves, *page.rects] if (c.get("fill") or c.get("stroke")) and is_bullet(c)]
    for dot in dots:
        free.append(Word("•", dot["x0"], dot["x1"] + 0.5, dot["top"] - 2, dot["bottom"] + 2, max(6.0, (dot["bottom"] - dot["top"]) * 2.6), "Arial", False, False, colour_hex(dot.get("non_stroking_color"))))
    lines = group_lines(free)

    # Underlines: a thin rule hugging the bottom of words is formatting, not a divider.
    keep_rules = []
    for rule in rules:
        if rule["dir"] == "h" and rule["thick"] <= 1.6:
            under = [w for line in lines for w in line.words if w.x0 >= rule["x0"] - 2 and w.x1 <= rule["x1"] + 2 and 0 <= rule["y"] - w.bottom + 0.2 * w.size <= 0.35 * w.size]
            if under and (rule["x1"] - rule["x0"]) <= sum(w.x1 - w.x0 for w in under) + 6 * len(under) + 4:
                for w in under:
                    w.underline = True
                continue
        keep_rules.append(rule)

    # Rules not belonging to a table: dividers in the flow (vertical ones can't be expressed outside a table).
    for rule in keep_rules:
        if rule["dir"] != "h":
            continue
        box = (rule["x0"], rule["y"] - rule["thick"] / 2, rule["x1"], rule["y"] + rule["thick"] / 2)
        if any(overlaps(box, t, 1.5) for t in table_boxes):
            continue
        if any(e.kind == "image" and inside(box, (e.x0, e.top, e.x1, e.bottom), 1) for e in elements):
            continue
        if rule["x1"] - rule["x0"] < 8:
            continue
        elements.append(Element("rule", box[0], box[1], box[2], box[3], data=rule))

    for line in lines:
        elements.append(Element("text", line.x0, line.top, line.x1, line.bottom, lines=[line]))

    # A form: frame edges and field lines everywhere, or a "drawing" that holds much of the page's text.
    stray_verticals = [r for r in keep_rules if r["dir"] == "v" and r["y1"] - r["y0"] > 10 and not any(overlaps((r["x"] - 1, r["y0"], r["x"] + 1, r["y1"]), t, 1) for t in table_boxes)]
    figure_words = max((e.data.get("words", 0) for e in elements if e.kind == "image"), default=0)
    complex_page = len(stray_verticals) >= 4 or figure_words > 0.3 * max(1, len(words))
    shading = [(b["box"], b["fill"]) for b in boxes if not any(inside(b["box"], t, 2) for t in table_boxes) and (b["box"][2] - b["box"][0]) > 10 and (b["box"][3] - b["box"][1]) > 4]
    model = PageModel(width, height, elements, shading, floating)
    model.complex = complex_page
    return model


def find_boxes(rules: list[dict]) -> list[tuple[float, float, float, float]]:
    """Rectangles outlined by rules: a top and bottom edge of the same extent, joined by left and right edges."""
    hs = [r for r in rules if r["dir"] == "h" and r["x1"] - r["x0"] > 12]
    vs = [r for r in rules if r["dir"] == "v" and r["y1"] - r["y0"] > 8]
    found = []
    for a in hs:
        for b in hs:
            if b["y"] - a["y"] < 8 or abs(a["x0"] - b["x0"]) > 2.5 or abs(a["x1"] - b["x1"]) > 2.5:
                continue
            height = b["y"] - a["y"]
            def side(x):
                return any(abs(v["x"] - x) < 2.5 and min(v["y1"], b["y"]) - max(v["y0"], a["y"]) > 0.8 * height for v in vs)
            if side(a["x0"]) and side(a["x1"]):
                found.append((min(a["x0"], b["x0"]), a["y"], max(a["x1"], b["x1"]), b["y"]))
                break
    # Innermost first; drop boxes that contain another found box (those are frames around groups).
    return [box for box in found if not any(other != box and inside(other, box, 0.5) for other in found)]


def repair_with_ocr(words: list[Word], ocr_lines: list[dict]) -> list[Word]:
    """Devanagari text layers are often broken (missing glyph maps, vowel signs cut off): such a line's text is
    taken from the OCR line over it, keeping the PDF's own font, size and position."""
    lines = group_lines(words)
    out: list[Word] = []
    for line in lines:
        text = line.text
        if not (DEVANAGARI.search(text) and BROKEN_TEXT.search(text)):
            out.extend(line.words)
            continue
        box = (line.x0, line.top, line.x1, line.bottom)
        hits = [o for o in ocr_lines if overlaps(box, tuple(o["box"]), 0) and DEVANAGARI.search(o["text"])]
        if not hits:
            out.extend(w for w in line.words if "\x00" not in w.text)
            continue
        hits.sort(key=lambda o: o["box"][0])
        first = line.words[0]
        out.append(Word(" ".join(o["text"] for o in hits), line.x0, line.x1, line.top, line.bottom, line.size, first.font, first.bold, first.italic, first.colour))
    return out


def cluster_boxes(boxes, pad):
    clusters = [list(b) for b in boxes]
    changed = True
    while changed:
        changed = False
        merged = []
        for box in clusters:
            for target in merged:
                if overlaps(box, target, pad):
                    target[0], target[1], target[2], target[3] = min(target[0], box[0]), min(target[1], box[1]), max(target[2], box[2]), max(target[3], box[3])
                    changed = True
                    break
            else:
                merged.append(box)
        clusters = merged
    return [tuple(c) for c in clusters]


def estimate_skew(grey: np.ndarray) -> float:
    """The page's tilt in degrees (projection profile: text rows are sharpest when the page is straight)."""
    import cv2

    small = cv2.resize(grey, (1000, max(1, int(grey.shape[0] * 1000 / grey.shape[1]))), interpolation=cv2.INTER_AREA)
    ink = (cv2.adaptiveThreshold(small, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY_INV, 31, 15) > 0).astype(np.float32)
    h, w = ink.shape

    def score(angle: float) -> float:
        matrix = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
        rows = cv2.warpAffine(ink, matrix, (w, h), flags=cv2.INTER_NEAREST).sum(axis=1)
        return float((np.diff(rows) ** 2).sum())

    coarse = max(np.arange(-4.0, 4.01, 0.25), key=score)
    return float(max(np.arange(coarse - 0.25, coarse + 0.26, 0.05), key=score))


def whiten(crop: Image.Image, paper: float) -> Image.Image:
    """A crop of a scan with its paper tone lifted to white (the picture, not the paper's grey or yellow)."""
    a = np.asarray(crop.convert("RGB"), dtype=np.float32)
    black = float(np.percentile(a.mean(axis=2), 1))
    if paper - black < 40:
        return crop
    out = np.clip((a - black) / (paper - black) * 255.0, 0, 255).astype(np.uint8)
    return Image.fromarray(out)


def analyse_scanned(pdf_doc, page_index: int, ocr_lines: list[dict], dpi: int = 200) -> PageModel:
    """A page that is a picture: straightened, its text from OCR lines; logos, photos, signatures, stamps and
    rules cropped from the scan (what's inked but isn't text), so nothing visual is lost and no picture is
    replaced by OCR's guess at text inside it."""
    import cv2

    page = pdf_doc[page_index]
    width, height = page.get_size()
    scale = dpi / 72
    image = page.render(scale=scale).to_pil().convert("RGB")
    grey = np.asarray(image.convert("L"), dtype=np.uint8)
    paper = float(np.percentile(grey, 92))

    # Straighten: the page and every OCR box turn together.
    angle = estimate_skew(grey)
    lines = [dict(line) for line in ocr_lines]
    if abs(angle) >= 0.15:
        h, w = grey.shape
        matrix = cv2.getRotationMatrix2D((w / 2, h / 2), angle, 1.0)
        rgb = cv2.warpAffine(np.asarray(image), matrix, (w, h), flags=cv2.INTER_CUBIC, borderValue=(int(paper),) * 3)
        image = Image.fromarray(rgb)
        grey = np.asarray(image.convert("L"), dtype=np.uint8)
        slope = abs(math.tan(math.radians(angle)))
        for line in lines:
            x0, top, x1, bottom = line["box"]
            cx, cy = (x0 + x1) / 2 * scale, (top + bottom) / 2 * scale
            nx = (matrix[0, 0] * cx + matrix[0, 1] * cy + matrix[0, 2]) / scale
            ny = (matrix[1, 0] * cx + matrix[1, 1] * cy + matrix[1, 2]) / scale
            bw = x1 - x0
            # A tilted line's upright box is taller than its text by the tilt across its width.
            bh = max((bottom - top) * 0.5, (bottom - top) - slope * bw)
            line["box"] = [nx - bw / 2, ny - bh / 2, nx + bw / 2, ny + bh / 2]

    ink = cv2.adaptiveThreshold(grey, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY_INV, 35, 15)
    # Colour that isn't grey ink (stamps, logos) counts even when light.
    hsv = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2HSV)
    ink |= ((hsv[..., 1] > 70) & (hsv[..., 2] > 40)).astype(np.uint8) * 255
    heights = [l["box"][3] - l["box"][1] for l in lines] or [10.0]
    line_height = float(np.median(heights))
    elements: list[Element] = []

    def text_cover(box) -> float:
        x0, top, x1, bottom = box
        area = max(1.0, (x1 - x0) * (bottom - top))
        return sum(area_of(intersection(box, tuple(l["box"]))) for l in lines) / area

    # Pictures: large inked shapes that text doesn't explain (a logo, a stamp, a signature, a photo). Text OCR
    # "read" inside them is part of the picture.
    joined = cv2.dilate(ink, cv2.getStructuringElement(cv2.MORPH_ELLIPSE, (int(5 * scale), int(5 * scale))))
    count, _, stats, _ = cv2.connectedComponentsWithStats(joined, connectivity=8)
    raw_count, raw_labels, raw_stats, _ = cv2.connectedComponentsWithStats(ink, connectivity=8)
    pictures = []
    for i in range(1, count):
        x, y, w, h, _ = stats[i]
        box = (x / scale, y / scale, (x + w) / scale, (y + h) / scale)
        if h / scale < 1.6 * line_height or w / scale < 1.5 * line_height:
            continue
        if w / scale > 0.9 * width and h / scale < 4 * line_height:
            continue  # a rule or a band, handled below
        # Text is many small marks (letters); a signature or a drawing is a few long strokes.
        pieces = raw_stats[1:]
        marks = int(((pieces[:, 0] >= x) & (pieces[:, 1] >= y) & (pieces[:, 0] + pieces[:, 2] <= x + w) & (pieces[:, 1] + pieces[:, 3] <= y + h) & (pieces[:, 4] > 3)).sum())
        stroke_like = marks < max(4, 0.25 * (w / scale) / line_height)
        if h / scale < 2.2 * line_height and not stroke_like:
            continue
        inside_lines = [l for l in lines if inside(tuple(l["box"]), box, 2)]
        # Text paragraphs join into big shapes too: those are mostly covered by confident OCR lines.
        confident = [l for l in inside_lines if l.get("confidence", 1) >= 0.6]
        if not stroke_like and text_cover(box) > 0.35 and len(confident) >= 2:
            continue
        pictures.append(box)
    # One picture per object, however many shapes it was found as.
    pictures = cluster_boxes(pictures, 2)
    for line in lines:
        x0, top, x1, bottom = line["box"]
        centre = ((x0 + x1) / 2, (top + bottom) / 2)
        line["drop"] = line.get("confidence", 1) < 0.5 or any(b[0] <= centre[0] <= b[2] and b[1] <= centre[1] <= b[3] for b in pictures)

    # Rules: long, thin horizontal runs that aren't text.
    text_mask = np.zeros_like(ink)
    for line in lines:
        if line["drop"]:
            continue
        x0, top, x1, bottom = line["box"]
        pad = 0.15 * (bottom - top)
        cv2.rectangle(text_mask, (int((x0 - pad) * scale), int((top - pad) * scale)), (int((x1 + pad) * scale), int((bottom + pad) * scale)), 255, -1)
    rest = cv2.bitwise_and(ink, cv2.bitwise_not(text_mask))
    horizontal = cv2.morphologyEx(rest, cv2.MORPH_OPEN, cv2.getStructuringElement(cv2.MORPH_RECT, (int(40 * scale), 1)))
    count, _, stats, _ = cv2.connectedComponentsWithStats(horizontal, connectivity=8)
    for i in range(1, count):
        x, y, w, h, _ = stats[i]
        box = (x / scale, y / scale, (x + w) / scale, (y + h) / scale)
        if h <= 2.2 * scale and not any(overlaps(box, b, 0) for b in pictures):
            elements.append(Element("rule", *box, data={"dir": "h", "x0": box[0], "x1": box[2], "y": (box[1] + box[3]) / 2, "thick": max(0.5, h / scale), "colour": "333333", "dashed": False}))
    for box in pictures:
        x0, top, x1, bottom = (int(box[0] * scale), int(box[1] * scale), int(box[2] * scale), int(box[3] * scale))
        crop = whiten(image.crop((x0, top, x1, bottom)), paper)
        elements.append(Element("image", *box, data={"png": png_bytes(crop)}))

    # Sizes from the straightened lines' heights, snapped to the page's few real sizes (a line without
    # descenders measures shorter than one with them, but it's the same type).
    kept = [l for l in lines if not l["drop"]]
    raw_sizes = sorted((l["box"][3] - l["box"][1]) * (0.62 if DEVANAGARI.search(l["text"]) else 0.74) for l in kept)
    clusters: list[list[float]] = []
    for value in raw_sizes:
        if clusters and value <= clusters[-1][0] * 1.22:
            clusters[-1].append(value)
        else:
            clusters.append([value])
    def snapped(value: float) -> float:
        for cluster in clusters:
            if cluster[0] <= value <= cluster[-1]:
                return float(np.median(cluster))
        return value

    words: list[Word] = []
    for line in kept:
        x0, top, x1, bottom = line["box"]
        devanagari = bool(DEVANAGARI.search(line["text"]))
        # A Devanagari line's box also holds its headstroke and vowel signs.
        size = max(5.0, round(snapped((bottom - top) * (0.62 if devanagari else 0.74)) * 2) / 2)
        words.append(Word(line["text"], x0, x1, top, bottom, size, line.get("font") or "Arial", bool(line.get("bold")), False, line.get("colour"), ocr=True))
    for line in group_lines(words):
        elements.append(Element("text", line.x0, line.top, line.x1, line.bottom, lines=[line]))
    model = PageModel(float(width), float(height), elements, [], [])
    model.complex = False
    return model


def area_of(box) -> float:
    return max(0.0, box[2] - box[0]) * max(0.0, box[3] - box[1])


# ------------------------------------------------------------------------------------------------ layout


def blocks_from(elements: list[Element]) -> list[Element]:
    """Text lines → paragraph blocks: lines close below each other with overlapping extents belong together.
    Short pieces that share their row with others (a form's labels, authors, "Ref … Date") stay on their own."""
    texts = sorted((e for e in elements if e.kind == "text"), key=lambda e: (e.top, e.x0))
    others = [e for e in elements if e.kind != "text"]
    lines_all = [e.lines[0] for e in texts]
    span = (max((l.x1 for l in lines_all), default=0) - min((l.x0 for l in lines_all), default=0)) or 1

    def tabular(line: Line) -> bool:
        if line.x1 - line.x0 > 0.35 * span:
            return False
        return any(o is not line and abs(o.baseline - line.baseline) < 2 and (o.x1 < line.x0 or o.x0 > line.x1) and o.x1 - o.x0 <= 0.35 * span for o in lines_all)

    blocks: list[Element] = []
    for element in texts:
        line = element.lines[0]
        if tabular(line):
            blocks.append(Element("text", line.x0, line.top, line.x1, line.bottom, lines=[line], data={"tabular": True}))
            continue
        best = None
        for block in reversed(blocks[-12:]):
            if block.data.get("tabular"):
                continue
            last = block.lines[-1]
            gap = line.top - last.bottom
            size = max(line.size, last.size)
            overlap = min(block.x1, line.x1) - max(block.x0, line.x0)
            if not (-0.3 * size <= gap <= 0.75 * size and overlap > 0.3 * min(line.x1 - line.x0, block.x1 - block.x0)):
                continue
            # The paragraph directly above. It continues unless this line starts a new one.
            if abs(line.size - last.size) > 0.12 * size:
                break
            rights = [l.x1 for l in block.lines]
            justified = len(block.lines) >= 2 and max(rights[:-1]) - min(rights[:-1]) < 0.6 * size
            if justified and last.x1 < max(rights) - 1.5 * size:
                break  # a justified paragraph's short last line ends it
            if len(block.lines) >= 2 and line.x0 > block.x0 + 0.6 * size and abs(last.x1 - block.x1) < 0.6 * size and line.x1 > block.x1 - 0.6 * size:
                break  # an indented first line after a full one starts the next paragraph
            if LIST_MARKER.match(line.text):
                break  # a list item starts its own paragraph
            if len(block.lines) >= 2:
                pitch = block.lines[-1].baseline - block.lines[-2].baseline
                if abs((line.baseline - last.baseline) - pitch) > 0.15 * pitch:
                    break  # set at a different pitch: not the same paragraph
            best = block
            break
        if best:
            best.lines.append(line)
            best.x0, best.x1, best.bottom = min(best.x0, line.x0), max(best.x1, line.x1), max(best.bottom, line.bottom)
        else:
            blocks.append(Element("text", line.x0, line.top, line.x1, line.bottom, lines=[line]))
    return blocks + others


LIST_MARKER = re.compile(r"^(?:[•▪◦‣·○■□●–—-]|\(?\d{1,3}[.)]|\(?[a-zA-Z][.)]|\(?[ivxIVX]{1,4}[.)])\s")


@dataclass
class Node:
    kind: str  # "flow" | "columns"
    x0: float
    top: float
    x1: float
    bottom: float
    items: list = field(default_factory=list)  # flow: Elements; columns: Nodes


def bounds(items):
    return min(i.x0 for i in items), min(i.top for i in items), max(i.x1 for i in items), max(i.bottom for i in items)


def xy_cut(items: list[Element], depth: int = 0) -> Node:
    """Recursive XY cut over blocks. A gutter that runs the full height splits columns; otherwise the region is cut
    across its widest horizontal gap. Leaves are flows read top to bottom."""
    x0, top, x1, bottom = bounds(items)
    if len(items) == 1 or depth > 12:
        return Node("flow", x0, top, x1, bottom, sorted(items, key=lambda e: (e.top, e.x0)))
    # Vertical gutters: x ranges no item covers.
    spans = sorted((i.x0, i.x1) for i in items if i.kind != "rule" or i.width < 0.6 * (x1 - x0))
    gutters = []
    reach = spans[0][1] if spans else x1
    for a, b in spans[1:]:
        if a - reach >= 6:
            gutters.append((reach, a))
        reach = max(reach, b)
    if gutters:
        # Cut at every gutter; full-width rules stay with whichever column they fall in.
        cuts = [(g[0] + g[1]) / 2 for g in gutters]
        groups: list[list[Element]] = [[] for _ in range(len(cuts) + 1)]
        spanning = []
        for item in items:
            centre = (item.x0 + item.x1) / 2
            if item.kind == "rule" and any(item.x0 < c < item.x1 for c in cuts):
                spanning.append(item)
                continue
            groups[sum(centre > c for c in cuts)].append(item)
        groups = [g for g in groups if g]
        if len(groups) > 1 and not spanning:
            return Node("columns", x0, top, x1, bottom, [xy_cut(g, depth + 1) for g in groups])
    # Horizontal cuts: every gap no item covers.
    ordered = sorted(items, key=lambda e: e.top)
    bands: list[list[Element]] = [[ordered[0]]]
    reach = ordered[0].bottom
    for item in ordered[1:]:
        if item.top >= reach - 0.5:
            bands.append([item])
        else:
            bands[-1].append(item)
        reach = max(reach, item.bottom)
    if len(bands) > 1:
        # Bands without full-width content (a title, an abstract, a wide figure) are joined back together, so a
        # region of running columns is cut down its gutter as a whole — not sliced wherever both columns
        # happen to have a gap at the same height (which would make the text read across the columns).
        region_width = x1 - x0
        spanning = [any(i.width > 0.6 * region_width for i in band) for band in bands]
        # Runs of consecutive narrow bands (indexes into `bands`), each run a candidate column region.
        runs: list[list[int]] = []
        for index, wide in enumerate(spanning):
            if runs and not wide and not spanning[runs[-1][-1]]:
                runs[-1].append(index)
            else:
                runs.append([index])

        # Only where the joined region really is running columns (each side holds several lines); otherwise
        # the bands stay apart (a "Ref … Date" row above an address is two rows, not two columns).
        def columns_of(items: list[Element]) -> bool:
            spans_ = sorted((i.x0, i.x1) for i in items if i.kind != "rule")
            if not spans_:
                return False
            reach_ = spans_[0][1]
            for a_, b_ in spans_[1:]:
                if a_ - reach_ >= 6:
                    left = [i for i in items if i.x1 <= a_ and i.kind == "text"]
                    right = [i for i in items if i.x0 >= a_ and i.kind == "text"]
                    return sum(len(i.lines) for i in left) >= 3 and sum(len(i.lines) for i in right) >= 3
                reach_ = max(reach_, b_)
            return False

        regrouped: list[list[Element]] = []
        for run in runs:
            group = [item for index in run for item in bands[index]]
            if len(run) > 1 and columns_of(group):
                regrouped.append(group)
            else:
                regrouped.extend(bands[index] for index in run)
        if len(regrouped) < len(bands):
            bands = regrouped
        children = [xy_cut(b, depth + 1) for b in bands]
        flat: list = []
        for child in children:
            if child.kind == "flow":
                flat.extend(child.items)
            else:
                flat.append(child)
        return Node("flow", x0, top, x1, bottom, flat)
    return Node("flow", x0, top, x1, bottom, sorted(items, key=lambda e: (e.top, e.x0)))


# ------------------------------------------------------------------------------------------------ writing


class Writer:
    def __init__(self):
        from docx import Document

        self.doc = Document()
        body = self.doc.element.body
        for child in list(body):
            if child.tag.endswith("}p") or child.tag.endswith("}tbl"):
                body.remove(child)
        self.body = body
        self.sect_template = body.find(qn("w:sectPr"))
        self.breaks_on_page = 0
        self.conflicts = 0
        self.drift = 0.0
        self.last_break = 0
        styles = self.doc.styles["Normal"]
        styles.font.name = "Arial"
        styles.font.size = _pt(10)
        styles.paragraph_format.space_after = _pt(0)
        styles.paragraph_format.space_before = _pt(0)

    # --- sections -----------------------------------------------------------------------------
    def section_xml(self, page: PageModel, margins, columns=None, continuous=False):
        from copy import deepcopy

        sect = deepcopy(self.sect_template)
        for child in list(sect):
            if child.tag in (qn("w:pgSz"), qn("w:pgMar"), qn("w:cols"), qn("w:type"), qn("w:docGrid")):
                sect.remove(child)
        if continuous:
            sect.insert(0, make("w:type", {"w:val": "continuous"}))
        w, h = page.width, page.height
        sect.append(make("w:pgSz", {"w:w": str(round(w * TWIPS)), "w:h": str(round(h * TWIPS)), **({"w:orient": "landscape"} if w > h else {})}))
        left, top, right, bottom = margins
        sect.append(make("w:pgMar", {"w:top": str(round(top * TWIPS)), "w:right": str(round(right * TWIPS)), "w:bottom": str(round(bottom * TWIPS)), "w:left": str(round(left * TWIPS)), "w:header": "0", "w:footer": "0", "w:gutter": "0"}))
        cols = make("w:cols", {"w:space": "0"})
        if columns:
            cols = make("w:cols", {"w:num": str(len(columns)), "w:equalWidth": "0", "w:space": "0"})
            for width, space in columns:
                cols.append(make("w:col", {"w:w": str(round(width * TWIPS)), "w:space": str(round(space * TWIPS))}))
        sect.append(cols)
        return sect

    def end_section(self, sect):
        """Ends the current section with `sect`: Word keeps a section's settings on its last paragraph."""
        last = self.body[-1] if len(self.body) else None
        if last is None or not last.tag.endswith("}p") or last.find("./" + qn("w:pPr") + "/" + qn("w:sectPr")) is not None or last.find("./" + qn("w:pPr") + "/" + qn("w:framePr")) is not None:
            last = self.spacer(0.1)
        ppr = last.find(qn("w:pPr"))
        if ppr is None:
            ppr = make("w:pPr")
            last.insert(0, ppr)
        ppr.append(sect)

    def finish(self, sect):
        old = self.body.find(qn("w:sectPr"))
        if old is not None:
            self.body.remove(old)
        self.body.append(sect)

    # --- primitives ---------------------------------------------------------------------------
    def paragraph(self, parent=None):
        p = make("w:p")
        (parent if parent is not None else self.body).append(p)
        return p

    def spacer(self, height: float, parent=None):
        """An empty paragraph exactly `height` points tall (vertical space before a table or after a section)."""
        p = self.paragraph(parent)
        ppr = p_props(p)
        ppr.append(make("w:spacing", {"w:before": "0", "w:after": "0", "w:line": str(max(2, round(height * TWIPS))), "w:lineRule": "exact"}))
        rpr = make("w:rPr")
        rpr.append(make("w:sz", {"w:val": "2"}))
        ppr.append(rpr)
        return p


def qn(tag: str) -> str:
    from docx.oxml.ns import qn as _qn

    return _qn(tag)


def make(tag: str, attrs: dict | None = None):
    from docx.oxml import OxmlElement

    element = OxmlElement(tag)
    for key, value in (attrs or {}).items():
        element.set(qn(key) if ":" in key else key, value)
    return element


def p_props(p):
    ppr = p.find(qn("w:pPr"))
    if ppr is None:
        ppr = make("w:pPr")
        p.insert(0, ppr)
    return ppr


# Word rejects a file whose properties are out of schema order (LibreOffice is forgiving, Word is not).
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr", "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents", "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr", "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"]


def reorder(element, order):
    rank = {name: i for i, name in enumerate(order)}
    children = list(element)
    children.sort(key=lambda child: rank.get(child.tag.split("}")[-1], len(order)))
    for child in children:
        element.remove(child)
        element.append(child)


def tidy(root):
    """Puts every pPr and tcPr in schema order (done once, after the document is built)."""
    for ppr in root.iter(qn("w:pPr")):
        reorder(ppr, PPR_ORDER)
    for tcpr in root.iter(qn("w:tcPr")):
        reorder(tcpr, TCPR_ORDER)


def _pt(value: float):
    from docx.shared import Pt

    return Pt(value)


def add_run(p, word_text: str, w: Word, space_before: bool, lang_ne: bool, squeeze: float = 0.0):
    r = make("w:r")
    rpr = make("w:rPr")
    fonts = {"w:ascii": w.font, "w:hAnsi": w.font, "w:eastAsia": w.font, "w:cs": w.font if DEVANAGARI.search(word_text) or "Devanagari" in w.font else w.font}
    rpr.append(make("w:rFonts", fonts))
    if w.bold:
        rpr.append(make("w:b"))
        rpr.append(make("w:bCs"))
    if w.italic:
        rpr.append(make("w:i"))
        rpr.append(make("w:iCs"))
    if w.colour and w.colour != "000000":
        rpr.append(make("w:color", {"w:val": w.colour}))
    if squeeze:
        rpr.append(make("w:spacing", {"w:val": str(round(squeeze * TWIPS))}))
    half = str(max(2, round(w.size * 2)))
    rpr.append(make("w:sz", {"w:val": half}))
    rpr.append(make("w:szCs", {"w:val": half}))
    if w.underline:
        rpr.append(make("w:u", {"w:val": "single"}))
    if lang_ne:
        rpr.append(make("w:lang", {"w:val": "en-US", "w:bidi": "ne-NP"}))
    r.append(rpr)
    t = make("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = CONTROL.sub("", (" " if space_before else "") + word_text)
    r.append(t)
    p.append(r)
    return r


def add_break(p):
    r = make("w:r")
    r.append(make("w:br"))
    p.append(r)


def add_tab(p):
    r = make("w:r")
    r.append(make("w:tab"))
    p.append(r)


# Font files with the same character widths as the fonts Word will use (Liberation and Arimo are metric twins).
FONT_FILES = {
    ("Arial", False, False): ["/System/Library/Fonts/Supplemental/Arial.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf", "/usr/share/fonts/truetype/msttcorefonts/Arial.ttf"],
    ("Arial", True, False): ["/System/Library/Fonts/Supplemental/Arial Bold.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf"],
    ("Arial", False, True): ["/System/Library/Fonts/Supplemental/Arial Italic.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf"],
    ("Arial", True, True): ["/System/Library/Fonts/Supplemental/Arial Bold Italic.ttf", "/usr/share/fonts/truetype/liberation/LiberationSans-BoldItalic.ttf"],
    ("Times New Roman", False, False): ["/System/Library/Fonts/Supplemental/Times New Roman.ttf", "/usr/share/fonts/truetype/liberation/LiberationSerif-Regular.ttf"],
    ("Times New Roman", True, False): ["/System/Library/Fonts/Supplemental/Times New Roman Bold.ttf", "/usr/share/fonts/truetype/liberation/LiberationSerif-Bold.ttf"],
    ("Times New Roman", False, True): ["/System/Library/Fonts/Supplemental/Times New Roman Italic.ttf", "/usr/share/fonts/truetype/liberation/LiberationSerif-Italic.ttf"],
    ("Times New Roman", True, True): ["/System/Library/Fonts/Supplemental/Times New Roman Bold Italic.ttf", "/usr/share/fonts/truetype/liberation/LiberationSerif-BoldItalic.ttf"],
    ("Courier New", False, False): ["/System/Library/Fonts/Supplemental/Courier New.ttf", "/usr/share/fonts/truetype/liberation/LiberationMono-Regular.ttf"],
    ("Courier New", True, False): ["/System/Library/Fonts/Supplemental/Courier New Bold.ttf", "/usr/share/fonts/truetype/liberation/LiberationMono-Bold.ttf"],
    ("Georgia", False, False): ["/System/Library/Fonts/Supplemental/Georgia.ttf"],
    ("Georgia", True, False): ["/System/Library/Fonts/Supplemental/Georgia Bold.ttf"],
    ("Georgia", False, True): ["/System/Library/Fonts/Supplemental/Georgia Italic.ttf"],
    ("Verdana", False, False): ["/System/Library/Fonts/Supplemental/Verdana.ttf"],
    ("Verdana", True, False): ["/System/Library/Fonts/Supplemental/Verdana Bold.ttf"],
}
_fonts: dict = {}


def text_width(text: str, w: Word) -> float | None:
    """How wide `text` is in the font Word will use, in points (None when no metric-compatible file is here)."""
    import os

    from PIL import ImageFont

    family = w.font if any(f == w.font for f, _, _ in FONT_FILES) else ("Times New Roman" if re.search(r"serif|roman|book|math|cambria|garamond|minion", w.font, re.I) and not re.search(r"sans", w.font, re.I) else "Arial")
    key = (family, w.bold, w.italic)
    if key not in _fonts:
        path = next((candidate for candidate in FONT_FILES.get(key, FONT_FILES.get((family, False, False), [])) if os.path.exists(candidate)), None)
        _fonts[key] = ImageFont.truetype(path, 100) if path else None
    font = _fonts[key]
    if font is None:
        return None
    return font.getlength(text) * w.size / 100


def line_squeeze(line: Line) -> float:
    """Character spacing (points, ≤ 0) that makes the line no wider than it is in the PDF, so it never wraps
    where the PDF didn't. Zero when it already fits or can't be measured."""
    if DEVANAGARI.search(line.text):
        return 0.0
    total = 0.0
    for index, word in enumerate(line.words):
        width = text_width(word.text + (" " if index < len(line.words) - 1 else ""), word)
        if width is None:
            return 0.0
        total += width
    # A point of room to spare: Word and LibreOffice break a line that is even a hair too wide.
    target = line.x1 - line.x0 - 1.0
    chars = max(1, len(line.text) - 1)
    if total <= target:
        return 0.0
    return max(-1.0, math.floor((target - total) / chars * TWIPS) / TWIPS)


def same_style(a: Word, b: Word):
    return a.font == b.font and a.bold == b.bold and a.italic == b.italic and a.colour == b.colour and abs(a.size - b.size) < 0.3 and a.underline == b.underline


def write_line_runs(p, line: Line, first_in_paragraph: bool):
    """A line's words as runs: consecutive words that look the same share a run; a line after the first is a
    soft line break, so the PDF's own line breaks are kept (and justified lines still fill their width)."""
    ne = any(DEVANAGARI.search(w.text) for w in line.words)
    if not first_in_paragraph:
        add_break(p)
    # A space only where the PDF leaves a gap: "text," with a bold word and a plain comma stays together.
    def spaced(previous: Word, word: Word) -> bool:
        return word.x0 - previous.x1 > 0.12 * min(previous.size, word.size)

    groups: list[list[Word]] = []
    for word in line.words:
        if groups and same_style(groups[-1][-1], word):
            groups[-1].append(word)
        else:
            groups.append([word])
    squeeze = line_squeeze(line)
    previous = None
    for group in groups:
        text = group[0].text
        for a, b in zip(group, group[1:]):
            text += (" " if spaced(a, b) else "") + b.text
        add_run(p, text, group[0], space_before=previous is not None and spaced(previous, group[0]), lang_ne=ne, squeeze=squeeze)
        previous = group[-1]


def alignment(block: Element, left: float, right: float) -> str:
    lines = block.lines
    if len(lines) == 1:
        line = lines[0]
        width = right - left
        centre = (line.x0 + line.x1) / 2
        if abs(centre - (left + right) / 2) < 0.03 * width and line.x0 - left > 0.12 * width:
            return "center"
        if right - line.x1 < 0.02 * width and line.x0 - left > 0.35 * width:
            return "right"
        return "left"
    lefts = [l.x0 for l in lines]
    rights = [l.x1 for l in lines]
    centres = [(l.x0 + l.x1) / 2 for l in lines]
    spread = lambda values: max(values) - min(values)  # noqa: E731
    size = statistics.median(l.size for l in lines)
    if spread(lefts[1:]) < 0.6 * size and spread(rights[:-1]) < 0.6 * size and len(lines) > 2:
        return "both"
    if spread(centres) < 0.8 * size and spread(lefts) > 1.5 * size:
        return "center"
    if spread(rights) < 0.6 * size and spread(lefts) > 0.6 * size and right - max(rights) < 0.5 * (right - left):
        return "right"
    if len(lines) == 2 and spread(lefts) < 0.6 * size and right - rights[0] < 0.8 * size and rights[1] < rights[0] - size:
        return "both"
    return "left"


class Flow:
    """Writes a flow of blocks top to bottom, turning each block's position into exact space before."""

    def __init__(self, writer: Writer, parent, left: float, right: float, cursor: float, page: PageModel):
        self.w = writer
        self.parent = parent
        self.left = left
        self.right = right
        self.cursor = cursor
        self.page = page

    def gap_to(self, top: float) -> float:
        gap = top - self.cursor
        if gap < 0:
            self.w.drift += -gap  # how far behind the PDF the flow has fallen
        if gap < -6:
            # Something here would have to sit above what's already placed: the page doesn't read as a flow.
            self.w.conflicts += 1
        return max(0.0, gap)

    def text(self, block: Element):
        lines = block.lines
        sizes = [l.size for l in lines]
        size = statistics.median(sizes)
        if all(w.ocr for l in lines for w in l.words):
            # Estimated sizes: one paragraph is one size.
            for l in lines:
                for w in l.words:
                    w.size = size
        if len(lines) > 1:
            pitch = statistics.median(b.baseline - a.baseline for a, b in zip(lines, lines[1:]))
            pitch = max(pitch, size * 0.95)
        else:
            pitch = size * 1.18
        top = lines[0].baseline - BASELINE_RATIO * pitch
        p = self.w.paragraph(self.parent)
        ppr = p_props(p)
        align = alignment(block, self.left, self.right)
        ppr.append(make("w:spacing", {"w:before": str(round(self.gap_to(top) * TWIPS)), "w:after": "0", "w:line": str(round(pitch * TWIPS)), "w:lineRule": "exact"}))
        indent_left = (min(l.x0 for l in lines[1:]) if len(lines) > 1 else lines[0].x0) - self.left
        first = lines[0].x0 - self.left - indent_left
        if align in ("center", "right"):
            # Only the side the text is set against is indented: the other keeps all the room, so a slightly wider
            # substitute font never wraps the line.
            indent_left, first = 0, 0
            right_indent = max(0.0, self.right - block.x1) if align == "right" else 0.0
            if align == "center":
                centre = (block.x0 + block.x1) / 2
                offset = centre - (self.left + self.right) / 2
                indent_left, right_indent = max(0.0, 2 * offset), max(0.0, -2 * offset)
            ind = {"w:left": str(round(max(0, indent_left) * TWIPS)), "w:right": str(round(right_indent * TWIPS))}
        else:
            right_indent = max(0.0, self.right - max(l.x1 for l in lines)) if align == "both" else 0.0
            ind = {"w:left": str(round(max(0.0, indent_left) * TWIPS)), "w:right": str(round(right_indent * TWIPS))}
            # A list item: marker, then the text at a hanging indent with a tab between them.
            if first < -1:
                ind["w:hanging"] = str(round(-first * TWIPS))
            elif first > 1:
                ind["w:firstLine"] = str(round(first * TWIPS))
        ppr.append(make("w:ind", ind))
        ppr.append(make("w:jc", {"w:val": align}))
        shade = self.shading_for(block)
        if shade:
            ppr.append(make("w:shd", {"w:val": "clear", "w:color": "auto", "w:fill": shade}))
        marker_gap = lines[0].words[1].x0 - lines[0].words[0].x1 if len(lines[0].words) > 1 else 0.0
        hanging_marker = len(lines[0].words) > 1 and bool(LIST_MARKER.match(lines[0].text)) and (first < -1 or (len(lines) == 1 and marker_gap > 0.4 * size))
        if hanging_marker and len(lines) == 1:
            # One-line item: hang the text at its own start so a later edit wraps under it.
            text_start = lines[0].words[1].x0 - self.left
            ind = {"w:left": str(round(max(0.0, text_start) * TWIPS)), "w:right": "0", "w:hanging": str(round(max(0.0, text_start - (lines[0].x0 - self.left)) * TWIPS))}
            for child in list(ppr):
                if child.tag == qn("w:ind"):
                    ppr.remove(child)
            ppr.append(make("w:ind", ind))
            indent_left = text_start
        for index, line in enumerate(lines):
            if index == 0 and hanging_marker:
                marker, rest = line.words[0], Line(line.words[1:])
                add_run(p, marker.text, marker, False, False)
                tabs = make("w:tabs")
                tabs.append(make("w:tab", {"w:val": "left", "w:pos": str(round(max(0.0, indent_left) * TWIPS))}))
                ppr.append(tabs)
                add_tab(p)
                write_line_runs(p, rest, True)
            else:
                write_line_runs(p, line, index == 0)
        self.cursor = top + pitch * len(lines)

    def shading_for(self, block: Element):
        for box, fill in self.page.shading:
            if inside((block.x0, block.top, block.x1, block.bottom), box, 3):
                return fill
        return None

    def tabbed(self, chunks: list[Element]):
        """One row of separate pieces of text ("Ref: …        Date: …"): one paragraph, tab stops where they start."""
        lines = [c.lines[0] for c in chunks]
        size = max(l.size for l in lines)
        pitch = size * 1.18
        baseline = statistics.median(l.baseline for l in lines)
        top = baseline - BASELINE_RATIO * pitch
        p = self.w.paragraph(self.parent)
        ppr = p_props(p)
        tabs = make("w:tabs")
        for line in lines[1:]:
            near_right = self.right - line.x1 < 3
            pos = line.x1 if near_right else line.x0
            tabs.append(make("w:tab", {"w:val": "right" if near_right else "left", "w:pos": str(round((pos - self.left) * TWIPS))}))
        ppr.append(tabs)
        ppr.append(make("w:spacing", {"w:before": str(round(self.gap_to(top) * TWIPS)), "w:after": "0", "w:line": str(round(pitch * TWIPS)), "w:lineRule": "exact"}))
        ppr.append(make("w:ind", {"w:left": str(round(max(0.0, lines[0].x0 - self.left) * TWIPS)), "w:right": "0"}))
        for index, line in enumerate(lines):
            if index:
                add_tab(p)
            write_line_runs(p, line, True)
        self.cursor = top + pitch

    def rule(self, element: Element):
        rule = element.data
        thick = max(0.25, min(6.0, rule["thick"]))
        # The rule paragraph is 1 pt tall; its bottom border is the rule.
        top = rule["y"] - thick / 2 - 1.0
        p = self.w.paragraph(self.parent)
        ppr = p_props(p)
        border = make("w:pBdr")
        border.append(make("w:bottom", {"w:val": "dashed" if rule.get("dashed") else "single", "w:sz": str(max(2, min(96, round(thick * 8)))), "w:space": "0", "w:color": rule.get("colour") or "000000"}))
        ppr.append(border)
        ppr.append(make("w:spacing", {"w:before": str(round(self.gap_to(top) * TWIPS)), "w:after": "0", "w:line": "20", "w:lineRule": "exact"}))
        ppr.append(make("w:ind", {"w:left": str(round(max(0.0, element.x0 - self.left) * TWIPS)), "w:right": str(round(max(0.0, self.right - element.x1) * TWIPS))}))
        rpr = make("w:rPr")
        rpr.append(make("w:sz", {"w:val": "2"}))
        ppr.append(rpr)
        self.cursor = top + 1.0 + thick

    def image(self, element: Element):
        gap = self.gap_to(element.top)
        p = self.w.paragraph(self.parent)
        ppr = p_props(p)
        ppr.append(make("w:spacing", {"w:before": str(round(gap * TWIPS)), "w:after": "0", "w:line": "240", "w:lineRule": "auto"}))
        ppr.append(make("w:ind", {"w:left": str(round(max(0.0, element.x0 - self.left) * TWIPS)), "w:right": "0"}))
        rpr = make("w:rPr")
        rpr.append(make("w:sz", {"w:val": "2"}))
        ppr.append(rpr)
        inline_picture(self.w.doc, p, element.data["png"], element.width, element.height)
        self.cursor = element.top + element.height + 0.5

    def table(self, element: Element):
        gap = self.gap_to(element.top)
        if gap >= 1.0:
            self.w.spacer(gap, self.parent)
        write_table(self.w, self.parent, element, self.left, self.page)
        self.cursor = element.bottom

    def side_by_side(self, node: Node):
        """Pieces next to each other that aren't running columns of text (a logo beside a name, an address beside
        a title): a borderless table, one cell per piece, each laid out by its own flow. A grid of single lines
        that line up row by row (labels and values) is rows of tab stops instead."""
        leaves = list(leaf_elements(node))
        if leaves and all(e.kind == "text" and len(e.lines) == 1 for e in leaves):
            rows: list[list[Element]] = []
            for leaf in sorted(leaves, key=lambda e: e.lines[0].baseline):
                if rows and abs(rows[-1][0].lines[0].baseline - leaf.lines[0].baseline) < 2.5:
                    rows[-1].append(leaf)
                else:
                    rows.append([leaf])
            if len(rows) < len(leaves):  # pieces really do share rows
                for row in rows:
                    row.sort(key=lambda e: e.x0)
                    if len(row) == 1:
                        self.text(row[0])
                    else:
                        self.tabbed(row)
                return
        gap = self.gap_to(node.top)
        if gap >= 1.0:
            self.w.spacer(gap, self.parent)
        columns = node.items
        edges = [self.left] + [(a.x1 + b.x0) / 2 for a, b in zip(columns, columns[1:])] + [max(self.right, columns[-1].x1)]
        tbl = new_table(self.parent, [edges[i + 1] - edges[i] for i in range(len(columns))], max(0.0, 0.0))
        row = make("w:tr")
        tbl.append(row)
        for index, column in enumerate(columns):
            cell = make("w:tc")
            tcpr = make("w:tcPr")
            tcpr.append(make("w:tcW", {"w:w": str(round((edges[index + 1] - edges[index]) * TWIPS)), "w:type": "dxa"}))
            cell.append(tcpr)
            row.append(cell)
            flow = Flow(self.w, cell, edges[index], edges[index + 1], node.top, self.page)
            flow.write(column)
            if not len(cell.findall(qn("w:p"))) or cell[-1].tag != qn("w:p"):
                self.w.paragraph(cell)
        self.cursor = node.bottom

    def write(self, node: Node):
        if node.kind == "columns":
            return self.side_by_side(node)
        items = node.items
        index = 0
        while index < len(items):
            item = items[index]
            if isinstance(item, Node):
                if item.kind == "columns" and all(isinstance(c, Node) and c.kind == "flow" and len(c.items) == 1 and isinstance(c.items[0], Element) and c.items[0].kind == "text" and len(c.items[0].lines) == 1 for c in item.items):
                    chunks = [c.items[0] for c in item.items]
                    if max(c.lines[0].baseline for c in chunks) - min(c.lines[0].baseline for c in chunks) < 2.5:
                        self.tabbed(chunks)
                        index += 1
                        continue
                self.write(item) if item.kind == "flow" else self.side_by_side(item)
            elif item.kind == "text":
                self.text(item)
            elif item.kind == "rule":
                self.rule(item)
            elif item.kind == "image":
                self.image(item)
            elif item.kind == "table":
                self.table(item)
            index += 1


def new_table(parent, widths: list[float], indent: float, borders: bool = False):
    tbl = make("w:tbl")
    tblpr = make("w:tblPr")
    tblpr.append(make("w:tblW", {"w:w": str(round(sum(widths) * TWIPS)), "w:type": "dxa"}))
    tblpr.append(make("w:tblInd", {"w:w": str(round(indent * TWIPS)), "w:type": "dxa"}))
    if not borders:
        b = make("w:tblBorders")
        for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
            b.append(make(f"w:{side}", {"w:val": "nil"}))
        tblpr.append(b)
    tblpr.append(make("w:tblLayout", {"w:type": "fixed"}))
    margins = make("w:tblCellMar")
    for side in ("top", "left", "bottom", "right"):
        margins.append(make(f"w:{side}", {"w:w": "0", "w:type": "dxa"}))
    tblpr.append(margins)
    tbl.append(tblpr)
    grid = make("w:tblGrid")
    for width in widths:
        grid.append(make("w:gridCol", {"w:w": str(round(width * TWIPS))}))
    tbl.append(grid)
    parent.append(tbl)
    return tbl


def write_table(writer: Writer, parent, element: Element, left: float, page: PageModel):
    """A ruled table: the grid from the cell edges, merged cells spanned, each side bordered only where the PDF
    draws it, cells shaded where the PDF fills them, text aligned as it sits in its cell."""
    cells = element.data["cells"]
    xs = sorted(snap({round(v, 1) for c in cells for v in (c[0], c[2])}))
    ys = sorted(snap({round(v, 1) for c in cells for v in (c[1], c[3])}))
    widths = [b - a for a, b in zip(xs, xs[1:])]
    words = element.data["words"]
    edges = table_edges(page, element)

    def left_margin(cell_box) -> float:
        mine = [w for w in words if inside(word_box(w), cell_box, 1.5)]
        if not mine:
            return 2.0
        lines = group_lines(mine)
        x0, x1 = min(l.x0 for l in lines), max(l.x1 for l in lines)
        if x1 < cell_box[2] - 6 and (x0 - cell_box[0]) > (cell_box[2] - x1) + 3:
            return 2.0  # right-aligned content: the cell keeps its full width, the paragraph sets the place
        return max(1.0, min(x0 - cell_box[0], 12))

    first = next((c for c in cells if abs(c[0] - xs[0]) < 2 and abs(c[1] - (snap({round(v, 1) for c2 in cells for v in (c2[1], c2[3])})[0])) < 2), None)
    # In this compatibility mode Word and LibreOffice measure the table's indent to the first cell's text.
    tbl = new_table(parent, widths, element.x0 - left + (left_margin(first) if first else 2.0), borders=False)
    grid: dict[tuple[int, int], tuple] = {}
    for c in cells:
        ci, cj = nearest(xs, c[0]), nearest(xs, c[2])
        ri, rj = nearest(ys, c[1]), nearest(ys, c[3])
        if cj > ci and rj > ri:
            grid[(ri, ci)] = (ri, rj, ci, cj, c)
    covered = {}
    for (ri, ci), (r0, r1, c0, c1, box) in grid.items():
        for r in range(r0, r1):
            for c in range(c0, c1):
                covered[(r, c)] = (r0, c0)
    for r in range(len(ys) - 1):
        tr = make("w:tr")
        trpr = make("w:trPr")
        # Borders add their width to a row: take it back so rows keep the PDF's pitch.
        border = max((e["thick"] for e in edges if e["dir"] == "h" and abs(e["y"] - ys[r + 1]) < 2.5), default=0.0)
        trpr.append(make("w:trHeight", {"w:val": str(round(max(4.0, ys[r + 1] - ys[r] - border) * TWIPS)), "w:hRule": "atLeast"}))
        tr.append(trpr)
        tbl.append(tr)
        c = 0
        while c < len(xs) - 1:
            owner = covered.get((r, c))
            if owner is None:
                tc = make("w:tc")
                tcpr = make("w:tcPr")
                tcpr.append(make("w:tcW", {"w:w": str(round(widths[c] * TWIPS)), "w:type": "dxa"}))
                tc.append(tcpr)
                tc.append(make("w:p"))
                tr.append(tc)
                c += 1
                continue
            r0, r1, c0, c1, box = grid[owner]
            span = c1 - c0
            tc = make("w:tc")
            tcpr = make("w:tcPr")
            tcpr.append(make("w:tcW", {"w:w": str(round(sum(widths[c0:c1]) * TWIPS)), "w:type": "dxa"}))
            if span > 1:
                tcpr.append(make("w:gridSpan", {"w:val": str(span)}))
            if r1 - r0 > 1:
                tcpr.append(make("w:vMerge", {"w:val": "restart"} if r == r0 else {}))
            borders = make("w:tcBorders")
            cell_box = (xs[c0], ys[r0], xs[c1], ys[r1])
            for side, present in cell_sides(cell_box, edges).items():
                thick, colour = present if present else (0, None)
                borders.append(make(f"w:{side}", {"w:val": "single", "w:sz": str(max(2, min(96, round(thick * 8)))), "w:space": "0", "w:color": colour} if present else {"w:val": "nil"}))
            tcpr.append(borders)
            fill = cell_fill(cell_box, page)
            if fill:
                tcpr.append(make("w:shd", {"w:val": "clear", "w:color": "auto", "w:fill": fill}))
            mine = [w for w in words if inside(word_box(w), cell_box, 1.5)] if r == r0 else []
            lines = group_lines(mine) if mine else []
            margin = left_margin(cell_box)
            if lines:
                first_top = min(l.top for l in lines)
                last_bottom = max(l.bottom for l in lines)
                room = cell_box[3] - cell_box[1]
                middle = abs((first_top + last_bottom) / 2 - (cell_box[1] + cell_box[3]) / 2) < 0.15 * room and first_top - cell_box[1] > 3
                if middle:
                    tcpr.append(make("w:vAlign", {"w:val": "center"}))
            mar = make("w:tcMar")
            mar.append(make("w:left", {"w:w": str(round(min(margin, 12) * TWIPS)), "w:type": "dxa"}))
            mar.append(make("w:right", {"w:w": str(round(2 * TWIPS)), "w:type": "dxa"}))
            tcpr.append(mar)
            tc.append(tcpr)
            if lines:
                for block in blocks_from([Element("text", l.x0, l.top, l.x1, l.bottom, lines=[l]) for l in lines]):
                    block_lines = sorted(block.lines, key=lambda l: l.top)
                    p = make("w:p")
                    tc.append(p)
                    ppr = p_props(p)
                    size = statistics.median(l.size for l in block_lines)
                    ppr.append(make("w:spacing", {"w:before": "0", "w:after": "0", "w:line": str(round(size * 1.15 * TWIPS)), "w:lineRule": "atLeast"}))
                    pad_left = min(l.x0 for l in block_lines) - cell_box[0]
                    pad_right = cell_box[2] - max(l.x1 for l in block_lines)
                    align = "right" if pad_right + 3 < pad_left and pad_left > 6 else "center" if abs(pad_left - pad_right) < 2.5 and pad_left > 6 else "left"
                    if align == "right":
                        ppr.append(make("w:ind", {"w:left": "0", "w:right": str(round(max(0.0, pad_right - 2) * TWIPS))}))
                    ppr.append(make("w:jc", {"w:val": align}))
                    for i, line in enumerate(block_lines):
                        write_line_runs(p, line, i == 0)
            else:
                tc.append(make("w:p"))
            tr.append(tc)
            c = c1
    return tbl


def snap(values: set[float], tolerance: float = 2.0) -> list[float]:
    out: list[float] = []
    for value in sorted(values):
        if out and value - out[-1] <= tolerance:
            continue
        out.append(value)
    return out


def nearest(values: list[float], value: float) -> int:
    return min(range(len(values)), key=lambda i: abs(values[i] - value))


def table_edges(page: PageModel, element: Element):
    return element.data.get("edges", [])


def cell_sides(box, edges):
    """Which sides of a cell the PDF draws (edges are thin rules near the table)."""
    x0, top, x1, bottom = box
    found = {"top": None, "left": None, "bottom": None, "right": None}
    for e in edges:
        if e["dir"] == "h":
            covers = min(e["x1"], x1) - max(e["x0"], x0) > 0.6 * (x1 - x0)
            if covers and abs(e["y"] - top) < 2.5:
                found["top"] = (e["thick"], e["colour"])
            if covers and abs(e["y"] - bottom) < 2.5:
                found["bottom"] = (e["thick"], e["colour"])
        else:
            covers = min(e["y1"], bottom) - max(e["y0"], top) > 0.6 * (bottom - top)
            if covers and abs(e["x"] - x0) < 2.5:
                found["left"] = (e["thick"], e["colour"])
            if covers and abs(e["x"] - x1) < 2.5:
                found["right"] = (e["thick"], e["colour"])
    return found


def cell_fill(box, page: PageModel):
    for shade_box, fill in page.table_fills if hasattr(page, "table_fills") else []:
        if inside(box, shade_box, 2.5) or (overlaps(box, shade_box, -1) and area(intersection(box, shade_box)) > 0.8 * area(box)):
            return fill
    return None


def area(b):
    return max(0.0, b[2] - b[0]) * max(0.0, b[3] - b[1])


def intersection(a, b):
    return (max(a[0], b[0]), max(a[1], b[1]), min(a[2], b[2]), min(a[3], b[3]))


_image_ids = [0]


def inline_picture(doc, p, data: bytes, width_pt: float, height_pt: float):
    """A picture in its paragraph, at its size on the page (python-docx builds the inline drawing)."""
    from docx.shared import Emu

    run = doc.add_paragraph().add_run()
    run.add_picture(io.BytesIO(data), width=Emu(int(width_pt * EMU_PER_PT)), height=Emu(int(height_pt * EMU_PER_PT)))
    holder = run._r.getparent()
    holder.getparent().remove(holder)
    p.append(run._r)


def floating_picture(doc, p, data: bytes, element: Element, behind: bool = True):
    """A picture anchored to the page at its exact place, behind the text (backgrounds, watermarks)."""
    from copy import deepcopy

    from docx.shared import Emu

    run = doc.add_paragraph().add_run()
    run.add_picture(io.BytesIO(data), width=Emu(int(element.width * EMU_PER_PT)), height=Emu(int(element.height * EMU_PER_PT)))
    holder = run._r.getparent()
    holder.getparent().remove(holder)
    inline = run._r.find(".//" + qn("wp:inline"))
    _image_ids[0] += 1
    anchor = make("wp:anchor", {"behindDoc": "1" if behind else "0", "distT": "0", "distB": "0", "distL": "0", "distR": "0", "simplePos": "0", "relativeHeight": str(_image_ids[0]), "locked": "0", "layoutInCell": "1", "allowOverlap": "1"})
    anchor.append(make("wp:simplePos", {"x": "0", "y": "0"}))
    h = make("wp:positionH", {"relativeFrom": "page"})
    off = make("wp:posOffset")
    off.text = str(int(element.x0 * EMU_PER_PT))
    h.append(off)
    v = make("wp:positionV", {"relativeFrom": "page"})
    off = make("wp:posOffset")
    off.text = str(int(element.top * EMU_PER_PT))
    v.append(off)
    anchor.append(h)
    anchor.append(v)
    for tag in ("wp:extent", "wp:effectExtent"):
        node = inline.find(qn(tag))
        if node is not None:
            anchor.append(deepcopy(node))
    anchor.append(make("wp:wrapNone"))
    for tag in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
        node = inline.find(qn(tag))
        if node is not None:
            anchor.append(deepcopy(node))
    inline.getparent().replace(inline, anchor)
    p.append(run._r)


# ------------------------------------------------------------------------------------------------ fixed layout


def textless_page_png(pdf_path: str, page_index: int, dpi: int = 150) -> bytes:
    """The page with its upright text taken out — its drawings, pictures and rotated text, exactly as printed —
    as the background a fixed-layout page's text sits on."""
    import pikepdf
    import pypdfium2 as pdfium

    def strip(owner, seen):
        if id(owner) in seen:
            return
        seen.add(id(owner))
        try:
            ops = pikepdf.parse_content_stream(owner)
        except Exception:  # noqa: BLE001
            return
        out, block, rotated = [], None, False
        rotation = [False]
        for item in ops:
            if not hasattr(item, "operator") or str(item.operator) == "INLINE IMAGE":
                (block if block is not None else out).append(item)  # inline images stay
                continue
            operands, op = item.operands, item.operator
            name = str(op)
            if block is None:
                if name == "q":
                    rotation.append(rotation[-1])
                elif name == "Q" and len(rotation) > 1:
                    rotation.pop()
                elif name == "cm" and len(operands) == 6 and (abs(float(operands[1])) > 1e-3 or abs(float(operands[2])) > 1e-3):
                    rotation[-1] = True
                if name == "BT":
                    block, rotated = [item], rotation[-1]
                    continue
                out.append(item)
            else:
                block.append(item)
                if name == "Tm" and len(operands) == 6 and (abs(float(operands[1])) > 1e-3 or abs(float(operands[2])) > 1e-3):
                    rotated = True
                if name == "ET":
                    if rotated:
                        out.extend(block)  # rotated text stays drawn: it can't be an upright Word line
                    block = None
        owner.write(pikepdf.unparse_content_stream(out))
        resources = owner.get("/Resources") if hasattr(owner, "get") else None
        for xobject in (resources.get("/XObject", {}) if resources is not None else {}).values():
            if xobject.get("/Subtype") == "/Form":
                strip(xobject, seen)

    with pikepdf.open(pdf_path) as pdf:
        page = pdf.pages[page_index]
        page.contents_coalesce()
        seen: set = set()
        strip(page.Contents, seen)
        for xobject in page.resources.get("/XObject", {}).values():
            if xobject.get("/Subtype") == "/Form":
                strip(xobject, seen)
        single = pikepdf.new()
        single.pages.append(page)
        buffer = io.BytesIO()
        single.save(buffer)
    doc = pdfium.PdfDocument(buffer.getvalue())
    image = doc[0].render(scale=dpi / 72).to_pil().convert("RGB")
    doc.close()
    out = io.BytesIO()
    image.save(out, format="PNG", optimize=True)
    return out.getvalue()


def reading_order(lines: list[Line]) -> list[Line]:
    """Lines in the order a reader takes them — column by column — so copying, searching and screen readers follow
    the text, even though each line is placed on its own."""
    if not lines:
        return []
    elements = [Element("text", l.x0, l.top, l.x1, l.bottom, lines=[l]) for l in lines]
    tree = xy_cut(blocks_from(elements))
    out: list[Line] = []

    def walk(node):
        for item in node.items:
            if isinstance(item, Node):
                walk(item)
            else:
                out.extend(item.lines)

    walk(tree)
    seen = {id(l) for l in out}
    return out + [l for l in lines if id(l) not in seen]


def write_fixed(writer: Writer, page: PageModel, background: bytes, lines: list[Line]):
    """A page Word can't express as a flow (a form, a dense layout): everything that isn't upright text is one
    picture behind the page, and every line of text is an editable paragraph framed at its exact place."""
    anchor = writer.spacer(1)
    floating_picture(writer.doc, anchor, background, Element("image", 0, 0, page.width, page.height), behind=True)
    for line in reading_order(lines):
        size = line.size
        pitch = size * 1.18
        p = writer.paragraph()
        ppr = p_props(p)
        y = line.baseline - BASELINE_RATIO * pitch
        # Wider than the text and free to grow downwards: a frame never clips its words.
        frame_width = (line.x1 - line.x0) * 1.15 + 8
        ppr.append(make("w:framePr", {"w:w": str(round(frame_width * TWIPS)), "w:h": str(round(pitch * TWIPS)), "w:hRule": "atLeast", "w:wrap": "around", "w:hAnchor": "page", "w:vAnchor": "page", "w:x": str(round(line.x0 * TWIPS)), "w:y": str(round(y * TWIPS))}))
        ppr.append(make("w:spacing", {"w:before": "0", "w:after": "0", "w:line": str(round(pitch * TWIPS)), "w:lineRule": "exact"}))
        ppr.append(make("w:ind", {"w:left": "0", "w:right": "0"}))
        write_line_runs(p, line, True)
    writer.spacer(1)


# ------------------------------------------------------------------------------------------------ pages


def write_page(writer: Writer, page: PageModel, is_last: bool):
    elements = blocks_from(page.elements)
    if not elements:
        sect = writer.section_xml(page, (36, 36, 36, 36))
        p = writer.paragraph()
        for f in page.floating:
            floating_picture(writer.doc, p, f.data["png"], f)
        return sect
    left = max(0.0, min(e.x0 for e in elements) - 0.5)
    right = min(page.width, max(e.x1 for e in elements) + 0.5)
    top = max(0.0, min(e.top for e in elements) - 2)
    bottom = max(0.0, min(18.0, page.height - max(e.bottom for e in elements) - 2))
    margins = (left, top, page.width - right, bottom)
    tree = xy_cut(elements)
    start_index = len(writer.body)
    writer.breaks_on_page = 0
    flow = Flow(writer, writer.body, left, right, top, page)
    write_tree(writer, flow, tree, page, margins, start_index)
    if page.floating:
        # Anchor backgrounds to the page's first paragraph.
        first = next((child for child in list(writer.body)[start_index:] if child.tag == qn("w:p")), None)
        if first is None:
            first = writer.paragraph()
        for f in page.floating:
            floating_picture(writer.doc, first, f.data["png"], f)
    return writer.section_xml(page, margins, continuous=writer.breaks_on_page > 0)


def write_tree(writer: Writer, flow: Flow, tree: Node, page: PageModel, margins, page_start: int):
    """Top-level flow; regions of running columns become Word section columns (continuous sections), so the
    columns read in order and stay side by side."""
    items = tree.items if tree.kind == "flow" else [tree]
    for item in items:
        if isinstance(item, Node) and item.kind == "columns" and is_article(item, page):
            if len(writer.body) > max(page_start, writer.last_break):
                # The part above the columns is a section of its own (a new page if it opened the page).
                writer.end_section(writer.section_xml(page, margins, continuous=writer.breaks_on_page > 0))
                writer.breaks_on_page += 1
                writer.last_break = len(writer.body)
            columns = item.items
            spans = []
            for index, column in enumerate(columns):
                start = flow.left if index == 0 else column.x0
                end = columns[index + 1].x0 if index + 1 < len(columns) else flow.right
                spans.append((max(1.0, column.x1 - start), max(0.0, end - column.x1) if index + 1 < len(columns) else 0.0))
            region_start = len(writer.body)
            for index, column in enumerate(columns):
                sub = Flow(writer, writer.body, flow.left if index == 0 else column.x0, column.x1, flow.cursor if index == 0 else item.top, page)
                sub.write(column)
                if index + 1 < len(columns):
                    last = writer.body[-1] if len(writer.body) > region_start and writer.body[-1].tag == qn("w:p") else writer.paragraph()
                    run = make("w:r")
                    run.append(make("w:br", {"w:type": "column"}))
                    last.append(run)
            writer.end_section(writer.section_xml(page, margins, columns=spans, continuous=writer.breaks_on_page > 0 or len(writer.body) > page_start))
            writer.breaks_on_page += 1
            writer.last_break = len(writer.body)
            flow.cursor = item.bottom
        elif isinstance(item, Node):
            flow.write(item if item.kind == "flow" else Node("flow", item.x0, item.top, item.x1, item.bottom, [item]))
        else:
            flow.write(Node("flow", item.x0, item.top, item.x1, item.bottom, [item]))


def is_article(node: Node, page: PageModel) -> bool:
    """Running columns of text (a paper, a newsletter) rather than pieces placed side by side.

    Off: Word and LibreOffice balance section columns differently (a picture near a column's foot jumps to the next
    column), so every column region is written as a borderless table instead, each column keeping its own content
    at its own place. Kept for the day section columns can be laid out predictably."""
    if True:
        return False
    if len(node.items) < 2:
        return False
    heights = [c.bottom - c.top for c in node.items]
    text_lines = [sum(len(e.lines) for e in leaf_elements(c) if e.kind == "text") for c in node.items]
    return min(heights) > 0.25 * page.height and min(text_lines) >= 8


def leaf_elements(node: Node):
    for item in node.items:
        if isinstance(item, Node):
            yield from leaf_elements(item)
        else:
            yield item


def attach_edges(model: PageModel, plumber_page):
    """Each table keeps the rules drawn around and inside it, for its cell borders; and the fills inside it."""
    rules, boxes = segments(plumber_page)
    model.table_fills = [(b["box"], b["fill"]) for b in boxes]
    for element in model.elements:
        if element.kind == "table":
            pad = 3
            box = (element.x0 - pad, element.top - pad, element.x1 + pad, element.bottom + pad)
            element.data["edges"] = [r for r in rules if (r["dir"] == "h" and box[1] <= r["y"] <= box[3] and r["x1"] > box[0] and r["x0"] < box[2]) or (r["dir"] == "v" and box[0] <= r["x"] <= box[2] and r["y1"] > box[1] and r["y0"] < box[3])]


def build_docx(pdf_path: str, out_path: str, pages: list[dict], progress=None) -> dict:
    """The Word file for these pages. Returns per-page notes (method, counts) for the caller's summary."""
    import pdfplumber
    import pypdfium2 as pdfium

    try:
        import pikepdf

        pike = pikepdf.open(pdf_path)
    except Exception:  # noqa: BLE001
        pike = None
    pdf_doc = pdfium.PdfDocument(pdf_path)
    writer = Writer()
    notes = []
    sections = []
    with pdfplumber.open(pdf_path) as plumber:
        for n, spec in enumerate(pages):
            index = int(spec["index"])
            plumber_page = plumber.pages[index]
            mode = spec.get("mode", "text")
            ocr_lines = spec.get("ocr") or None
            if mode != "ocr":
                # A scan with a hidden text layer (a "searchable" PDF): the picture is the page. It is treated as a
                # scan whose lines come from that layer, so the text isn't shown twice.
                page_area = float(plumber_page.width) * float(plumber_page.height)
                if any((im["x1"] - im["x0"]) * (im["bottom"] - im["top"]) > 0.8 * page_area for im in plumber_page.images):
                    layer_words, _ = read_words(plumber_page)
                    layer = [{"text": l.text, "box": [l.x0, l.top, l.x1, l.bottom], "size": l.size, "font": l.words[0].font, "bold": l.words[0].bold} for l in group_lines(layer_words)]
                    broken = any(DEVANAGARI.search(l["text"]) and BROKEN_TEXT.search(l["text"]) for l in layer)
                    if layer:
                        ocr_lines = ocr_lines if (broken and ocr_lines) else layer
                        mode = "ocr"
            if mode == "ocr" and ocr_lines is not None:
                model = analyse_scanned(pdf_doc, index, ocr_lines)
                model.table_fills = []
            else:
                model = analyse_digital(plumber_page, index, pdf_doc, pike, ocr_lines)
                attach_edges(model, plumber_page)
            start = len(writer.body)
            writer.conflicts = 0
            writer.drift = 0.0
            fixed = mode != "ocr" and (spec.get("layout") == "fixed" or model.complex)
            if not fixed:
                sect = write_page(writer, model, n == len(pages) - 1)
                # The flow couldn't hold the page's layout (things would have to overlap): start the page again, fixed.
                # …or it would end up far below where the PDF has it (and spill onto the next page).
                if mode != "ocr" and (writer.conflicts >= 3 or writer.drift > 14) and spec.get("layout") != "flow":
                    for child in list(writer.body)[start:]:
                        if child.tag != qn("w:sectPr"):
                            writer.body.remove(child)
                    fixed = True
            if fixed:
                words, _ = read_words(plumber_page)
                write_fixed(writer, model, textless_page_png(pdf_path, index), group_lines(words, split=0.6))
                sect = writer.section_xml(model, (0, 0, 0, 0))
                mode = "fixed"
            if n < len(pages) - 1:
                writer.end_section(sect)
            else:
                writer.finish(sect)
            notes.append({"page": index + 1, "method": mode, "pictures": sum(1 for e in model.elements if e.kind == "image") + len(model.floating), "tables": sum(1 for e in model.elements if e.kind == "table"), "rules": sum(1 for e in model.elements if e.kind == "rule")})
            if progress:
                progress(n + 1, len(pages))
    tidy(writer.body)
    writer.doc.save(out_path)
    pdf_doc.close()
    if pike is not None:
        pike.close()
    return {"pages": notes}
