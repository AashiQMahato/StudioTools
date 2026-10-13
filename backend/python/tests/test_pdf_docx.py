"""
PDF → Word: the analysis pieces on their own, and whole conversions of the test PDFs checked for what a reader
would notice — the right pages, the text, the pictures, tables and rules, and a valid Word file.
(The visual comparison through LibreOffice is scripts/docx_visual_qa.py.)
"""

from __future__ import annotations

import sys
import zipfile
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "image_service"))

import pdf_docx  # noqa: E402
from pdf_docx import Line, Word, blocks_from, build_docx, find_boxes, font_info, group_lines, merge_rules  # noqa: E402

FIXTURES = Path(__file__).parent / "docx_fixtures"


def word(text, x0, top, size=10.0, bold=False):
    return Word(text, x0, x0 + len(text) * size * 0.5, top, top + size, size, "Arial", bold, False, None)


@pytest.mark.parametrize(
    "name,expected",
    [
        ("AAAAAA+Helvetica-Bold", ("Arial", True, False)),
        ("NimbusRomNo9L-ReguItal", ("Times New Roman", False, True)),
        ("NimbusRomNo9L-Medi", ("Times New Roman", True, False)),
        ("HelveticaNeueLTStd-BdIt", ("Arial", True, True)),
        ("TimesNewRomanPS-BoldMT", ("Times New Roman", True, False)),
        ("CMBX10", ("Times New Roman", True, False)),
        ("JAAAAA+KohinoorDevanagari-Regular", ("Kohinoor Devanagari", False, False)),
    ],
)
def test_font_names_map_to_word_families_and_styles(name, expected):
    assert font_info(name) == expected


def test_words_split_into_pieces_at_wide_gaps_not_at_spaces():
    lines = group_lines([word("Ref:", 50, 100), word("ST/114", 75, 100), word("Date:", 400, 100)])
    assert [l.text for l in lines] == ["Ref: ST/114", "Date:"]


def test_dashes_merge_into_one_dashed_rule():
    pieces = [{"dir": "h", "x0": 10 + i * 4, "x1": 12 + i * 4, "y": 50, "thick": 0.7, "colour": "999999"} for i in range(20)]
    rules = merge_rules(pieces)
    assert len(rules) == 1 and rules[0]["dashed"] and rules[0]["x1"] == pytest.approx(88)


def test_four_rules_around_text_make_a_box():
    rules = merge_rules(
        [
            {"dir": "h", "x0": 50, "x1": 250, "y": 100, "thick": 1, "colour": "000000"},
            {"dir": "h", "x0": 50, "x1": 250, "y": 160, "thick": 1, "colour": "000000"},
            {"dir": "v", "y0": 100, "y1": 160, "x": 50, "thick": 1, "colour": "000000"},
            {"dir": "v", "y0": 100, "y1": 160, "x": 250, "thick": 1, "colour": "000000"},
        ]
    )
    assert find_boxes(rules) == [(50, 100, 250, 160)]


def test_list_items_and_pitch_changes_start_new_paragraphs():
    lines = [
        Line([word("• First item that wraps", 40, 100)]),
        Line([word("onto a second line", 40, 109)]),
        Line([word("• Second item", 40, 121)]),
        Line([word("Body text after the list.", 40, 140)]),
    ]
    elements = [pdf_docx.Element("text", l.x0, l.top, l.x1, l.bottom, lines=[l]) for l in lines]
    blocks = [b for b in blocks_from(elements) if b.kind == "text"]
    assert [len(b.lines) for b in blocks] == [2, 1, 1]


def docx_text(path: Path) -> str:
    import re

    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    return " ".join(re.findall(r"<w:t[^>]*>([^<]*)</w:t>", xml))


def test_letter_keeps_text_pictures_rules_and_page_size(tmp_path):
    out = tmp_path / "letter.docx"
    notes = build_docx(str(FIXTURES / "letterhead.pdf"), str(out), [{"index": 0, "mode": "text"}])
    page = notes["pages"][0]
    assert page["method"] == "text" and page["pictures"] == 3 and page["rules"] >= 2
    text = docx_text(out)
    for phrase in ("Studio Tools Pvt. Ltd.", "Subject: Confirmation of Service Agreement", "Managing Director", "Page 1 of 1"):
        assert phrase in text
    from docx import Document

    section = Document(str(out)).sections[0]
    assert section.page_width.pt == pytest.approx(595, abs=1) and section.page_height.pt == pytest.approx(842, abs=1)


def test_invoice_table_is_a_real_table_with_its_cells(tmp_path):
    out = tmp_path / "invoice.docx"
    notes = build_docx(str(FIXTURES / "invoice.pdf"), str(out), [{"index": 0, "mode": "text"}])
    assert notes["pages"][0]["tables"] >= 3  # the line items and the two address boxes
    from docx import Document

    tables = Document(str(out)).tables
    items = next(t for t in tables if len(t.rows) == 5)
    assert items.cell(1, 1).text.strip() == "Document scanning (A4, duplex)"
    assert items.cell(4, 4).text.strip() == "38,000.00"


def test_two_column_page_reads_column_by_column(tmp_path):
    out = tmp_path / "columns.docx"
    build_docx(str(FIXTURES / "columns.pdf"), str(out), [{"index": 0, "mode": "text"}])
    text = docx_text(out)
    # Column by column, as a reader goes: the first column (with its picture) before the second, which begins with
    # the picture's caption (Chrome set it at the top of the second column) — not interleaved line by line.
    assert text.index("1. Introduction") < text.index("2. Method") < text.index("Figure 1.") < text.index("3. Results") < text.index("4. Conclusion")


def test_scan_gets_ocr_text_and_cropped_pictures_not_ocr_noise(tmp_path):
    out = tmp_path / "scan.docx"
    # OCR lines as the Node pipeline sends them (points); one is a misread signature.
    lines = [
        {"text": "Subject: Confirmation of Service Agreement", "box": [67, 240, 385, 256], "confidence": 0.98},
        {"text": "Yours faithfully,", "box": [66, 482, 147, 495], "confidence": 0.97},
        {"text": "WMMW", "box": [70, 535, 200, 560], "confidence": 0.31},
    ]
    notes = build_docx(str(FIXTURES / "scanned-letter.pdf"), str(out), [{"index": 0, "mode": "ocr", "ocr": lines}])
    text = docx_text(out)
    assert "Subject: Confirmation of Service Agreement" in text and "WMMW" not in text
    assert notes["pages"][0]["pictures"] >= 2  # the logo and the signature/stamp, cropped from the scan


def test_dense_form_falls_back_to_an_exact_layout(tmp_path):
    w9 = FIXTURES / ".cache" / "w9.pdf"
    if not w9.exists():
        pytest.skip("the W-9 test form isn't downloaded (scripts/docx_visual_qa.py fetches nothing; add it to .cache)")
    out = tmp_path / "w9.docx"
    notes = build_docx(str(w9), str(out), [{"index": 0, "mode": "text"}])
    assert notes["pages"][0]["method"] == "fixed"
    assert "Request for Taxpayer" in docx_text(out)
